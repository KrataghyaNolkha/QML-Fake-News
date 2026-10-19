from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "response_to_reviewers.docx"


def add_response(doc: Document, label: str, response: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(response)


def main() -> None:
    doc = Document()
    doc.core_properties.author = "Anonymous"
    doc.core_properties.last_modified_by = "Anonymous"
    doc.core_properties.title = "Response to Reviewers"
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Response to Reviewers")
    doc.add_paragraph("Paper 687 | MPCIT 2026")
    doc.add_paragraph(
        "We thank the reviewers for their careful assessment. We revised the manuscript and regenerated the primary results from the full Facebook fact-check dataset. The responses below describe changes made in the revised manuscript and identify the scope of experiments explicitly."
    )

    doc.add_heading("Reviewer 1", level=1)
    add_response(doc, "Comment 1. Condition-number scales differ between tables.",
                 "We audited the diagnostic implementation and standardized the reported definition. We center and symmetrize the Gram matrix, add a diagonal ridge of 10^-8, and report the largest positive eigenvalue divided by the smallest positive eigenvalue. The regenerated run gives values on the 10^8–10^9 scale. We could not reproduce the earlier 30–45 values under this definition and removed them rather than presenting them as comparable results.")
    add_response(doc, "Comment 2. Table II contains negative spectral gaps.",
                 "We made descending eigenvalue ordering explicit and recomputed the results. The reported gap is lambda_1 minus lambda_2 after centering and ridge stabilization; the regenerated values are nonnegative. The implementation rejects invalid nonsquare or nonfinite matrices.")
    add_response(doc, "Comment 3. Regression coefficients conflict.",
                 "We regenerated one standardized OLS model from 24 distinct model settings, with accuracy as the response and alignment, log10(condition number), and spectral gap as predictors. The standardized coefficients are 0.431, -0.443, and -0.348, respectively; R-squared is 0.448. The corresponding predictor p-values are 0.067, 0.410, and 0.467. We removed conflicting coefficient sets and describe this small-sample model as exploratory.")
    add_response(doc, "Comment 4. Per-seed support-vector counts were missing.",
                 "We include the regenerated ten-seed table with each seed, accuracy, and the two solver-specific counts. The table defines the classical count as the sum of RBF SVC support vectors and the quantum count as distinct Pegasos active training examples. We omit a ratio because these counts are defined by different optimization procedures.")
    add_response(doc, "Comment 5. The single-run and aggregate ratios differ.",
                 "We withdrew both the earlier single-run ratio of 25.34 and the ten-seed mean ratio of 15.61. They divide counts produced by different solvers and therefore do not estimate a common compression or efficiency quantity. The revised table retains the raw per-seed counts with explicit solver-specific definitions; it reports no count ratio, inferential test on those counts, or computational-efficiency claim.")

    doc.add_heading("Reviewer 5", level=1)
    add_response(doc, "Comment 1. The geometry analysis has only 19 configurations.",
                 "The earlier table's 19 rows came from the standard preset: four feature-map comparisons, nine depth-sweep rows, and six PCA/qubit rows. We reran the full preset, which produced 25 sweep rows (four feature-map, 12 depth-sweep, and nine PCA/qubit rows). One Pauli PCA=2, qubits=2, reps=2, 200-iteration setting appeared in both the feature-map and PCA/qubit sweeps, so it is retained once in configuration-level analysis. The analysis therefore uses 24 distinct feature-map, PCA, qubit, depth, and iteration settings. We report the small sample explicitly and treat associations as exploratory, not causal or out-of-sample prediction.")
    add_response(doc, "Comment 2. Condition-number correlations conflict.",
                 "We traced the earlier draft values to its 25-row sweep table. Recomputing log10(condition number) against accuracy on all 25 rows reproduces Pearson r = -0.521 and Spearman rho = -0.517 (n=25); applying the draft's absolute z-score > 2 rule reproduces r = -0.672 and rho = -0.600 (n=23). The 25 rows include the same Pauli-XYZ, PCA=2, 2-qubit, 2-repetition, 200-iteration setting twice, once from each sweep. Removing that duplicate changes the all-setting result to r = -0.516 and rho = -0.519 (n=24), and the z-score-excluded result to r = -0.667 and rho = -0.606 (n=22). Thus the discrepancy came from counting one repeated configuration and recomputing the outlier rule on the changed sample; the earlier figures are reproducible but are superseded by the deduplicated analysis. The machine-readable reconciliation records both calculations. All settings share one train/test split, so we treat these small-sample associations as descriptive rather than independent-sample evidence.")
    add_response(doc, "Comment 3. Support-vector counts may not be comparable.",
                 "We now state exactly what each count measures. For the classical RBF SVC (gamma='scale', C=1), the count is the sum of scikit-learn's per-class support-vector counts. For quantum Pauli-XYZ Pegasos, it is the number of distinct training examples that entered the active set after a margin-violating update; the ten-seed run uses lambda=0.02 and 100 sampled updates. Both pipelines use seed-matched splits and two PCA components, but the classical features are standardized while the quantum inputs are MinMax-scaled to [0, pi]. Thus the models are matched by split and PCA dimension, while their input scaling and solvers differ. The manuscript reports the raw counts descriptively, removes the count ratio and count significance test, and makes no memory, inference-time, or compression claim. The run manifest records these settings. Longer-iteration sensitivity runs are not claimed as completed.")
    add_response(doc, "Comment 4. The main geometry study is limited to low qubit counts.",
                 "The full geometry sweep tested 2, 3, and 4 qubits, but its coverage is unbalanced. The four feature-map comparisons and 12 depth-sweep settings held qubit count at 2 to isolate those factors at fixed width. The narrower Pauli-XYZ PCA/qubit tradeoff tested 3 settings each at 2, 3, and 4 qubits across PCA dimensions 2, 3, and 4. After deduplication, the 24 settings comprise 18 at 2 qubits, 3 at 3 qubits, and 3 at 4 qubits; geometry diagnostics use n=400. We now state that this is not a balanced qubit-scaling study and limit conclusions to these settings. We make no claim about qubit counts above 4.")
    add_response(doc, "Comment 5. Cross-dataset validation does not replicate the full geometry analysis.",
                 "We revised the manuscript to state that cross-dataset generalization has not been established. The breast-cancer experiment is a separate classification sanity check; it does not reproduce the fake-news task, preprocessing, or geometry protocol and is not used to validate the geometry findings. WELFake and LIAR results are excluded from manuscript tables and claims because their saved outputs contain conflicting runs and incomplete provenance. We removed the earlier wording that implied the breast-cancer check ruled out a dataset-specific result. A matched, provenance-complete geometry evaluation across multiple fake-news datasets remains future work.")
    add_response(doc, "Comment 6. Noise results are simulator-only.",
                 "We identify the AerSimulator and modeled depolarizing gate errors. The finite-shot stress test uses 512 shots, a stratified 60-example training subset, a 30-example kernel-statistics subset, and the full held-out split for accuracy. The manuscript reports the nonmonotonic accuracy values alongside the decreasing mean kernel similarity and makes no hardware-robustness claim.")
    add_response(doc, "Comment 7. Correlations are at risk of causal interpretation.",
                 "We replaced causal or predictive phrasing with exploratory association language and report the sample size, p-values, and confidence intervals. The multivariate coefficients are not statistically significant at 0.05.")
    add_response(doc, "Comment 8. Classical models outperform the quantum model on other representations.",
                 "We identify the representation for each comparison. The quantum result is compared directly with the PCA(2) RBF baseline; higher full-feature and engineered-feature classical results are reported separately. We do not claim that the quantum models outperform the strongest classical model overall.")
    add_response(doc, "Comment 9. Reproducibility details are incomplete.",
                 "We expanded the methods with the retained-row and label rules; categorical encoding; stratified split and seed; train-fitted PCA and scaler order; padding/truncation for mismatched PCA and qubit dimensions; exact Qiskit feature-map families, entanglement choices, and ideal fidelity-kernel calculation; Pegasos update rule, lambda, and iteration counts; classical model hyperparameters and metric conventions; and the Aer noise channel, gate sets, error probabilities, shots, subsets, and transpiler optimization level. The run manifest now records these settings together with software versions. We also disclose that the PCA(2) five-fold check reuses PCA fitted on the outer training split, so it is not a fully fold-isolated preprocessing estimate.")
    add_response(doc, "Comment 10. Novelty is not clearly positioned.",
                 "We revised the related work to acknowledge prior fake-news QML directly, including Tian (2023), which evaluates QKNN/QSVM/VQC, and Khalil et al. (2025), which studies PegasosQSVM with PCA and Z/ZZ feature maps on BuzzFace and TruthSeeker. We therefore withdraw any implication that this paper is the first quantum fake-news classifier or proposes a new classifier, feature map, or Pegasos algorithm. We distinguish the present contribution as a bounded empirical diagnostic study: we report alignment, spectral summaries, and conditioning with accuracy across 24 distinct settings for structured Facebook fact-check metadata, alongside a ten-seed matched PCA(2) RBF comparison and a modeled finite-shot stress test. We added a comparison table and relevant references. We explicitly limit novelty to this task-specific empirical synthesis; we do not claim a new metric, causal mechanism, cross-dataset generalization, or quantum advantage.")

    doc.add_heading("Reviewer 6", level=1)
    add_response(doc, "Comment 1. Add implementation detail.",
                 "The experimental setup now states the data filters and label mapping, categorical encoding, train/test split, train-only PCA/scaling, feature-map circuit constructions and entanglement, exact fidelity kernel, Pegasos update and settings, baseline hyperparameters, geometry definitions, and simulator noise gate model. The run manifest provides the machine-readable configuration and package versions; the response does not infer undocumented circuit constants.")
    add_response(doc, "Comment 2. Standardize Table I.",
                 "The report generator now applies a consistent 8-point table font, centered cell alignment, repeated header rows, header shading, and scientific notation for large values. The regenerated manuscript applies yellow highlighting to the revised text.")
    add_response(doc, "Comment 3. Add recent related work and align the title with results.",
                 "We added prior quantum fake-news studies (Tian, 2023; Khalil et al., 2025) and explain how their classifier-performance focus differs from this paper's geometry-diagnostic focus. We also position our exploratory measurements alongside established quantum feature-space, kernel-concentration, and generalization work (Havlicek et al., 2019; Thanasilp et al., 2024; Gil-Fuster et al., 2024). The comparison table states both the overlap and the bounded distinction. The title remains descriptive of that contribution, and the abstract and conclusion do not claim quantum advantage.")

    doc.add_heading("Resubmission checks", level=1)
    doc.add_paragraph(
        "The revised manuscript is anonymized in its document properties and has no author-identifying title-page fields. All revised manuscript text is highlighted yellow. Tables and figures are regenerated from the completed full-dataset run."
    )
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
